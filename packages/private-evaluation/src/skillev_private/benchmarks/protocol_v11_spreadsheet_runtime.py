"""Protocol 11 read-only workbook inspection layered over the proven V10 sandbox/OJ."""

from __future__ import annotations

from dataclasses import asdict, dataclass, field, replace
from pathlib import Path

from skillev.benchmarks.protocol_v11_action import protocol_v11_action_contract
from skillev.contracts import JsonValue, normalize_json
from skillev.rollout import RolloutSessionBundle, RolloutTask
from skillev.runtime import ActionKind, BudgetVector, EnvironmentObservation, StructuredAction

from .protocol_v10_spreadsheet_runtime import IsolatedSpreadsheetWorkspace
from .spreadsheet_public_view import (
    SpreadsheetPromptContract,
    SpreadsheetPromptSetting,
    build_spreadsheet_public_view,
    public_cell_from_openpyxl,
    public_cell_value,
)


@dataclass(frozen=True, slots=True)
class WorkbookState:
    file_size: int
    modified_ns: int
    loadable: bool
    sheets: tuple[str, ...]


def inspect_workbook_state(path: Path) -> WorkbookState:
    from openpyxl import load_workbook

    stat = path.stat()
    try:
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            sheets = tuple(workbook.sheetnames)
        finally:
            workbook.close()
        return WorkbookState(stat.st_size, stat.st_mtime_ns, True, sheets)
    except (OSError, ValueError):
        return WorkbookState(stat.st_size, stat.st_mtime_ns, False, ())


@dataclass(slots=True)
class SpreadsheetInspectionRuntime:
    workbook_path: Path

    def _workbook(self) -> object:
        from openpyxl import load_workbook

        return load_workbook(self.workbook_path, read_only=True, data_only=False)

    def list_sheets(self) -> dict[str, JsonValue]:
        workbook = self._workbook()
        try:
            names = list(workbook.sheetnames)  # type: ignore[attr-defined]
        finally:
            workbook.close()  # type: ignore[attr-defined]
        return _public_object({"sheets": names})

    def preview(self, arguments: dict[str, JsonValue]) -> dict[str, JsonValue]:
        from openpyxl.utils.cell import range_boundaries

        sheet_name = arguments.get("sheet")
        if not isinstance(sheet_name, str) or not sheet_name:
            raise ValueError("spreadsheet preview requires a sheet")
        maximum_rows = _bounded_integer(arguments.get("max_rows", 20), maximum=50)
        maximum_columns = _bounded_integer(arguments.get("max_columns", 12), maximum=30)
        workbook = self._workbook()
        try:
            if sheet_name not in workbook.sheetnames:  # type: ignore[attr-defined]
                raise KeyError("unknown spreadsheet sheet")
            sheet = workbook[sheet_name]  # type: ignore[index]
            range_value = arguments.get("range")
            if range_value is None:
                min_col, min_row = 1, 1
            elif isinstance(range_value, str):
                raw_min_col, raw_min_row, _, _ = range_boundaries(range_value)
                if raw_min_col is None or raw_min_row is None:
                    raise ValueError("spreadsheet preview range is incomplete")
                min_col, min_row = raw_min_col, raw_min_row
            else:
                raise ValueError("spreadsheet preview range must be text or null")
            rows = [
                [asdict(public_cell_from_openpyxl(cell)) for cell in row]
                for row in sheet.iter_rows(
                    min_row=min_row,
                    max_row=min_row + maximum_rows - 1,
                    min_col=min_col,
                    max_col=min_col + maximum_columns - 1,
                )
            ]
            return _public_object(
                {
                    "sheet": sheet_name,
                    "start_row": min_row,
                    "start_column": min_col,
                    "rows": rows,
                }
            )
        finally:
            workbook.close()  # type: ignore[attr-defined]

    def describe_range(self, arguments: dict[str, JsonValue]) -> dict[str, JsonValue]:
        preview = self.preview({**arguments, "max_rows": 50, "max_columns": 30})
        rows = preview["rows"]
        assert isinstance(rows, list)
        cells = [
            cell for row in rows if isinstance(row, list) for cell in row if isinstance(cell, dict)
        ]
        return _public_object(
            {
                "sheet": preview["sheet"],
                "nonempty_cells": sum(cell.get("value") is not None for cell in cells),
                "formula_cells": sum(cell.get("formula") is not None for cell in cells),
                "data_type_counts": {
                    str(kind): sum(cell.get("data_type") == kind for cell in cells)
                    for kind in sorted(
                        {
                            str(cell["data_type"])
                            for cell in cells
                            if cell.get("data_type") is not None
                        }
                    )
                },
                "preview": rows,
                "truncated": len(rows) >= 50,
            }
        )

    def find(self, arguments: dict[str, JsonValue]) -> dict[str, JsonValue]:
        query = arguments.get("query")
        sheet_filter = arguments.get("sheet")
        start_after = arguments.get("start_after")
        limit = _bounded_integer(arguments.get("limit", 50), maximum=100)
        if not isinstance(query, str) or not query.strip():
            raise ValueError("spreadsheet find requires text")
        if sheet_filter is not None and not isinstance(sheet_filter, str):
            raise ValueError("spreadsheet find sheet must be text or null")
        if start_after is not None and not isinstance(start_after, str):
            raise ValueError("spreadsheet find cursor must be text or null")
        needle = query.casefold()
        workbook = self._workbook()
        matches: list[dict[str, JsonValue]] = []
        try:
            for sheet in workbook.worksheets:  # type: ignore[attr-defined]
                if sheet_filter is not None and sheet.title != sheet_filter:
                    continue
                for row in sheet.iter_rows(
                    max_row=min(sheet.max_row, 2_000), max_col=min(sheet.max_column, 100)
                ):
                    for cell in row:
                        cursor = f"{sheet.title}!{cell.coordinate}"
                        if start_after is not None and cursor <= start_after:
                            continue
                        if cell.value is not None and needle in str(cell.value).casefold():
                            matches.append(
                                {
                                    "sheet": sheet.title,
                                    "cell": cell.coordinate,
                                    "value": public_cell_value(cell.value),
                                }
                            )
                            if len(matches) == limit:
                                return _public_object(
                                    {
                                        "matches": matches,
                                        "truncated": True,
                                        "next_start_after": cursor,
                                    }
                                )
        finally:
            workbook.close()  # type: ignore[attr-defined]
        return _public_object({"matches": matches, "truncated": False, "next_start_after": None})


def _bounded_integer(value: JsonValue, *, maximum: int) -> int:
    if type(value) is not int or not 1 <= value <= maximum:
        raise ValueError(f"spreadsheet bound must be an integer in [1, {maximum}]")
    return value


def _public_object(value: object) -> dict[str, JsonValue]:
    normalized = normalize_json(value)
    if not isinstance(normalized, dict):
        raise TypeError("spreadsheet public observation must be an object")
    return normalized


@dataclass(slots=True)
class ProtocolV11SpreadsheetWorkspace:
    inner: IsolatedSpreadsheetWorkspace
    inspector: SpreadsheetInspectionRuntime = field(init=False)
    _last_step: int = field(default=0, init=False)

    def __post_init__(self) -> None:
        self.inspector = SpreadsheetInspectionRuntime(self.inner.directory / "input.xlsx")

    @property
    def environment_id(self) -> str:
        return self.inner.environment_id

    @property
    def task_family(self) -> str:
        return self.inner.task_family

    async def execute(self, action: StructuredAction, *, step_index: int) -> EnvironmentObservation:
        if step_index <= self._last_step:
            raise ValueError("spreadsheet steps must be strictly increasing")
        self._last_step = step_index
        if action.kind is not ActionKind.TOOL or action.resource_id != "spreadsheet":
            return await self.inner.execute(action, step_index=step_index)
        if not isinstance(action.arguments, dict):
            return _tool_error("spreadsheet arguments must be an object")
        try:
            if action.name == "list_sheets":
                value = self.inspector.list_sheets()
            elif action.name == "preview":
                value = self.inspector.preview(action.arguments)
            elif action.name == "describe_range":
                value = self.inspector.describe_range(action.arguments)
            elif action.name == "find":
                value = self.inspector.find(action.arguments)
            elif action.name == "execute":
                before = inspect_workbook_state(self.inspector.workbook_path)
                result = await self.inner.execute(action, step_index=step_index)
                after = inspect_workbook_state(self.inspector.workbook_path)
                public = _public_object(result.public_value)
                public = _public_object(
                    {
                        **public,
                        "workbook_changed": (before.file_size, before.modified_ns)
                        != (after.file_size, after.modified_ns),
                        "workbook_loadable": after.loadable,
                        "post_execution_sheets": list(after.sheets),
                    }
                )
                return EnvironmentObservation(
                    public_value=public,
                    observation_status=result.observation_status,
                    invoked_skill_ids=result.invoked_skill_ids,
                    budget_usage=result.budget_usage,
                )
            else:
                return _tool_error("unsupported spreadsheet action")
        except (KeyError, ValueError) as exc:
            return _tool_error(str(exc))
        return EnvironmentObservation(
            public_value=value,
            observation_status="success",
            budget_usage=BudgetVector(tool_calls=1),
        )

    def validate_completion(self, submission: JsonValue) -> bool:
        return self.inner.validate_completion(submission)

    async def grade(self, task_id: str, submitted_workbook: JsonValue) -> object:
        return await self.inner.grade(task_id, submitted_workbook)

    async def close(self) -> None:
        await self.inner.close()


def bind_protocol_v11_spreadsheet_session(
    task: RolloutTask,
    session: RolloutSessionBundle,
) -> tuple[RolloutTask, RolloutSessionBundle]:
    """Upgrade one proven V10 workbook/OJ route to the frozen V11 public contract."""

    if not isinstance(session.environment, IsolatedSpreadsheetWorkspace):
        raise TypeError("Protocol 11 spreadsheet binding requires the isolated V10 workspace")
    context = task.public_context if isinstance(task.public_context, dict) else {}
    contract = SpreadsheetPromptContract(
        setting=SpreadsheetPromptSetting.MULTI_ROUND_ROWS_AND_EXECUTION,
        expose_instruction_type=True,
        expose_answer_position=True,
    )
    workbook = session.environment.directory / "input.xlsx"
    view = build_spreadsheet_public_view(
        workbook_path=workbook,
        instruction=task.query,
        instruction_type=_optional_public_text(context, "instruction_type", "instruction-type"),
        answer_position=_optional_public_text(context, "answer_position", "answer-position"),
        contract=contract,
    )
    surface, budget = protocol_v11_action_contract("spreadsheetbench")
    tools = tuple(sorted(f"{tool.resource_id}:{tool.name}" for tool in surface.tools))
    upgraded_task = replace(
        task,
        context_id=f"{task.context_id}/spreadsheet-public-inspection-v11",
        available_tools=tools,
        public_context=normalize_json(asdict(view)),
        action_surface=surface,
        budget_profile=budget,
    )
    workspace = ProtocolV11SpreadsheetWorkspace(session.environment)
    upgraded_session = replace(session, environment=workspace, cleanup=workspace.close)
    return upgraded_task, upgraded_session


def _optional_public_text(context: dict[str, JsonValue], *names: str) -> str | None:
    for name in names:
        value = context.get(name)
        if isinstance(value, str) and value.strip():
            return value
    return None


def _tool_error(message: str) -> EnvironmentObservation:
    return EnvironmentObservation(
        public_value={"error": message},
        observation_status="tool_error",
        budget_usage=BudgetVector(tool_calls=1),
    )


__all__ = [
    "ProtocolV11SpreadsheetWorkspace",
    "SpreadsheetInspectionRuntime",
    "WorkbookState",
    "bind_protocol_v11_spreadsheet_session",
    "inspect_workbook_state",
]
