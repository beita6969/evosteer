"""Bounded model-visible SpreadsheetBench view derived only from input workbooks."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime, time
from enum import StrEnum
from pathlib import Path
from typing import Any

from skillev.contracts import JsonValue

_MAXIMUM_PUBLIC_CELL_TEXT_CHARACTERS = 256


class SpreadsheetPromptSetting(StrEnum):
    SINGLE_ROUND = "single-round"
    MULTI_ROUND_ROWS = "multi-round-rows"
    MULTI_ROUND_ROWS_AND_EXECUTION = "multi-round-rows-and-execution"


@dataclass(frozen=True, slots=True)
class SpreadsheetPromptContract:
    setting: SpreadsheetPromptSetting
    expose_instruction_type: bool
    expose_answer_position: bool
    preview_rows: int = 20
    preview_columns: int = 12
    maximum_total_cells: int = 600

    def __post_init__(self) -> None:
        if min(self.preview_rows, self.preview_columns, self.maximum_total_cells) <= 0:
            raise ValueError("spreadsheet preview limits must be positive")


@dataclass(frozen=True, slots=True)
class SpreadsheetSheetPreview:
    sheet_name: str
    max_row: int
    max_column: int
    first_nonempty_rows: tuple[tuple[PublicCell, ...], ...]
    formula_cell_count: int | None
    merged_range_count: int
    preview_truncated: bool


@dataclass(frozen=True, slots=True)
class SpreadsheetPublicTaskView:
    instruction: str
    workbook_filename: str
    sheets: tuple[SpreadsheetSheetPreview, ...]
    instruction_type: str | None
    answer_position: str | None
    output_contract: str


@dataclass(frozen=True, slots=True)
class PublicCell:
    coordinate: str
    value: JsonValue
    data_type: str
    formula: str | None
    number_format: str | None
    style_name: str | None


def public_cell_value(value: object) -> JsonValue:
    if value is None or type(value) in {str, int, float, bool}:
        if isinstance(value, str) and len(value) > _MAXIMUM_PUBLIC_CELL_TEXT_CHARACTERS:
            return value[: _MAXIMUM_PUBLIC_CELL_TEXT_CHARACTERS - 3] + "..."
        return value  # type: ignore[return-value]
    if isinstance(value, datetime | date | time):
        return value.isoformat()
    return str(value)


def public_cell_from_openpyxl(cell: Any) -> PublicCell:
    value = cell.value
    style = cell.style if hasattr(cell, "style") else None
    return PublicCell(
        coordinate=str(cell.coordinate),
        value=public_cell_value(value),
        data_type=str(cell.data_type),
        formula=(
            str(public_cell_value(value))
            if isinstance(value, str) and value.startswith("=")
            else None
        ),
        number_format=(str(cell.number_format) if cell.number_format else None),
        style_name=(str(style) if style else None),
    )


def build_spreadsheet_public_view(
    *,
    workbook_path: Path,
    instruction: str,
    instruction_type: str | None,
    answer_position: str | None,
    contract: SpreadsheetPromptContract,
) -> SpreadsheetPublicTaskView:
    """Read only the supplied input workbook under fixed scan limits."""

    from openpyxl import load_workbook

    if not workbook_path.is_file() or not instruction.strip():
        raise ValueError("spreadsheet public view requires an input workbook and instruction")
    workbook = load_workbook(workbook_path, read_only=False, data_only=False)
    previews: list[SpreadsheetSheetPreview] = []
    visible_cells_remaining = contract.maximum_total_cells
    try:
        for sheet in workbook.worksheets:
            row_limit = min(sheet.max_row, contract.preview_rows)
            column_limit = min(sheet.max_column, contract.preview_columns)
            rows: list[tuple[PublicCell, ...]] = []
            formula_count = 0
            scanned = 0
            truncated = sheet.max_row * sheet.max_column > visible_cells_remaining
            for raw_row in sheet.iter_rows(
                min_row=1,
                max_row=row_limit,
                max_col=column_limit,
                values_only=False,
            ):
                rendered = tuple(public_cell_from_openpyxl(cell) for cell in raw_row)
                if any(value.value is not None for value in rendered):
                    visible = rendered[:visible_cells_remaining]
                    if visible:
                        rows.append(visible)
                        visible_cells_remaining -= len(visible)
                    if len(visible) != len(rendered) or visible_cells_remaining == 0:
                        truncated = True
                        break
            for raw_row in sheet.iter_rows(
                min_row=1,
                max_row=min(sheet.max_row, contract.maximum_total_cells),
                max_col=min(sheet.max_column, contract.maximum_total_cells),
                values_only=False,
            ):
                for cell in raw_row:
                    if scanned >= contract.maximum_total_cells:
                        break
                    scanned += 1
                    if isinstance(cell.value, str) and cell.value.startswith("="):
                        formula_count += 1
                if scanned >= contract.maximum_total_cells:
                    break
            previews.append(
                SpreadsheetSheetPreview(
                    sheet_name=sheet.title,
                    max_row=sheet.max_row,
                    max_column=sheet.max_column,
                    first_nonempty_rows=tuple(rows),
                    formula_cell_count=None if truncated else formula_count,
                    merged_range_count=len(sheet.merged_cells.ranges),
                    preview_truncated=truncated,
                )
            )
    finally:
        workbook.close()
    return SpreadsheetPublicTaskView(
        instruction=instruction,
        workbook_filename="input.xlsx",
        sheets=tuple(previews),
        instruction_type=instruction_type if contract.expose_instruction_type else None,
        answer_position=answer_position if contract.expose_answer_position else None,
        output_contract="Edit WORKBOOK_PATH in place and submit.",
    )


__all__ = [
    "PublicCell",
    "SpreadsheetPromptContract",
    "SpreadsheetPromptSetting",
    "SpreadsheetPublicTaskView",
    "SpreadsheetSheetPreview",
    "build_spreadsheet_public_view",
    "public_cell_from_openpyxl",
    "public_cell_value",
]
