from pathlib import Path

from openpyxl import Workbook
from skillev_private.benchmarks.spreadsheet_public_view import (
    SpreadsheetPromptContract,
    SpreadsheetPromptSetting,
    build_spreadsheet_public_view,
)


def _workbook(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Summary"
    sheet["A1"] = "Revenue"
    sheet["B1"] = 10
    sheet["B2"] = "=B1*2"
    sheet.merge_cells("C1:D1")
    workbook.save(path)
    workbook.close()


def test_public_view_is_bounded_and_contract_controls_metadata(tmp_path: Path) -> None:
    path = tmp_path / "input.xlsx"
    _workbook(path)
    view = build_spreadsheet_public_view(
        workbook_path=path,
        instruction="Update the workbook.",
        instruction_type="formula",
        answer_position="Summary!B2",
        contract=SpreadsheetPromptContract(
            SpreadsheetPromptSetting.MULTI_ROUND_ROWS_AND_EXECUTION,
            expose_instruction_type=True,
            expose_answer_position=False,
            preview_rows=2,
            preview_columns=4,
        ),
    )
    assert view.workbook_filename == "input.xlsx"
    assert view.instruction_type == "formula"
    assert view.answer_position is None
    assert view.sheets[0].formula_cell_count == 1
    assert view.sheets[0].merged_range_count == 1


def test_large_view_marks_formula_count_unknown(tmp_path: Path) -> None:
    path = tmp_path / "input.xlsx"
    _workbook(path)
    view = build_spreadsheet_public_view(
        workbook_path=path,
        instruction="Update.",
        instruction_type=None,
        answer_position=None,
        contract=SpreadsheetPromptContract(
            SpreadsheetPromptSetting.MULTI_ROUND_ROWS,
            False,
            False,
            maximum_total_cells=1,
        ),
    )
    assert view.sheets[0].preview_truncated
    assert view.sheets[0].formula_cell_count is None


def test_public_view_caps_cell_text_and_total_visible_cells(tmp_path: Path) -> None:
    path = tmp_path / "input.xlsx"
    workbook = Workbook()
    first = workbook.active
    first["A1"] = "x" * 1_000
    second = workbook.create_sheet("Second")
    second["A1"] = "visible"
    workbook.save(path)
    workbook.close()

    view = build_spreadsheet_public_view(
        workbook_path=path,
        instruction="Update.",
        instruction_type=None,
        answer_position=None,
        contract=SpreadsheetPromptContract(
            SpreadsheetPromptSetting.MULTI_ROUND_ROWS,
            False,
            False,
            maximum_total_cells=1,
        ),
    )
    value = view.sheets[0].first_nonempty_rows[0][0].value
    assert isinstance(value, str)
    assert len(value) == 256
    assert not view.sheets[1].first_nonempty_rows
